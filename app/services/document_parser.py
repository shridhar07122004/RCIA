from __future__ import annotations

import hashlib
import re
from pathlib import Path


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}


class DocumentParseError(ValueError):
    pass


def validate_upload(filename: str, content_type: str | None, data: bytes, max_size: int) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError("Unsupported file type. Upload a PDF or DOCX document.")
    if len(data) == 0:
        raise DocumentParseError("The uploaded document is empty.")
    if len(data) > max_size:
        raise DocumentParseError(f"The document exceeds the {max_size // (1024 * 1024)} MB upload limit.")
    expected = {".pdf": "application/pdf", ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
    if suffix in expected and content_type and content_type not in {expected[suffix], "application/octet-stream"}:
        raise DocumentParseError("The file content type does not match its extension.")
    if suffix == ".pdf" and not data.startswith(b"%PDF"):
        raise DocumentParseError("This file is not a readable PDF or may be corrupted.")
    if suffix == ".docx" and not data.startswith(b"PK"):
        raise DocumentParseError("This file is not a readable DOCX or may be corrupted.")
    return suffix


def parse_document(path: Path, suffix: str) -> str:
    try:
        if suffix == ".pdf":
            try:
                import fitz
            except ImportError as exc:
                raise DocumentParseError("PDF support is unavailable. Install PyMuPDF with: pip install PyMuPDF") from exc
            document = fitz.open(path)
            if document.is_encrypted:
                raise DocumentParseError("This PDF is encrypted and cannot be processed.")
            text = "\n".join(page.get_text("text") for page in document)
            document.close()
            if not text.strip():
                raise DocumentParseError("No selectable text was found. Scanned PDFs need OCR, which is not enabled in this version.")
            return text
        if suffix == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise DocumentParseError("DOCX support is unavailable. Install python-docx with: pip install python-docx") from exc
            document = Document(path)
            text_parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    text_parts.append(" | ".join(cell.text for cell in row.cells))
            text = "\n".join(text_parts)
            if not text.strip():
                raise DocumentParseError("The DOCX contains no readable text.")
            return text
        text = path.read_text(encoding="utf-8", errors="replace")
        if not text.strip():
            raise DocumentParseError("The document contains no readable text.")
        return text
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError("The document could not be read safely.") from exc


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def clean_text(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()
